"""
Builds the project writeup in two formats from one source (scripts/writeup_content.py):

    project_writeup.pdf            (repo root -- required location per the project brief)
    outputs/project_writeup.docx   (Word version of the same document)

Inputs are what project_code.py produces: outputs/metrics_summary.json and the PNGs in
outputs/figures. Run project_code.py first, then:

    python scripts/build_writeup.py            # both formats
    python scripts/build_writeup.py --pdf      # PDF only
    python scripts/build_writeup.py --docx     # Word only

If img.png (team logo) is present in the repo root it is placed small in the page header;
otherwise the header is text-only.
"""
import argparse
import io
import json
from pathlib import Path

import numpy as np
from PIL import Image

from writeup_content import TITLE, RUNNING_HEADER, build_blocks

ROOT = Path(__file__).resolve().parent.parent
OUTPUTS_DIR = ROOT / "outputs"
FIGURES_DIR = OUTPUTS_DIR / "figures"
METRICS_PATH = OUTPUTS_DIR / "metrics_summary.json"
LOGO_PATH = ROOT / "img.png"
PDF_PATH = ROOT / "project_writeup.pdf"
DOCX_PATH = OUTPUTS_DIR / "project_writeup.docx"

NAVY = (31, 58, 94)
BLUE = (31, 111, 178)
RED = (184, 55, 48)
GRAY = (100, 100, 100)
LGRAY = (235, 237, 240)
ROW_FILL = (248, 249, 250)
BLACK = (25, 25, 25)
CALLOUT_COLORS = {"red": RED, "blue": BLUE}

MARGIN_MM = 18
LOGO_H_MM = 8


def load_logo():
    """img.png with its light, fake-transparent background trimmed off; None if absent."""
    if not LOGO_PATH.exists():
        return None
    a = np.asarray(Image.open(LOGO_PATH).convert("RGB")).astype(int)
    # foreground = anything dark or clearly colored; background = light low-saturation pixels
    fg = (a.min(axis=2) < 200) | ((a.max(axis=2) - a.min(axis=2)) > 40)
    ys, xs = np.where(fg)
    a[~fg] = 255
    return Image.fromarray(a.astype("uint8")).crop((xs.min(), ys.min(), xs.max() + 1, ys.max() + 1))


def is_bold_row(name):
    return "Final" in name or ("LightGBM" in name and "blend" not in name.lower())


# ---------------------------------------------------------------------------
# PDF (fpdf2)
# ---------------------------------------------------------------------------
def build_pdf(blocks, logo, path):
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            top = 6
            text_x = self.l_margin
            if logo is not None:
                self.image(logo, x=self.l_margin, y=top, h=LOGO_H_MM)
                text_x += LOGO_H_MM * logo.width / logo.height + 3
            if self.page_no() == 1:
                self.set_y(top + LOGO_H_MM + 4)
                return
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(*GRAY)
            text_y = top + LOGO_H_MM / 2 - 3
            self.set_xy(text_x, text_y)
            self.cell(0, 6, RUNNING_HEADER, align="L")
            self.set_xy(self.l_margin, text_y)
            self.cell(0, 6, f"{self.page_no()}", align="R", new_x="LMARGIN", new_y="NEXT")
            self.set_draw_color(*LGRAY)
            self.set_line_width(0.3)
            self.line(self.l_margin, top + LOGO_H_MM + 1.5, self.w - self.r_margin, top + LOGO_H_MM + 1.5)
            self.set_y(top + LOGO_H_MM + 5)

    pdf = PDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(MARGIN_MM, 20, MARGIN_MM)
    text_w = pdf.w - pdf.l_margin - pdf.r_margin

    def h1(text):
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 15)
        pdf.set_text_color(*NAVY)
        pdf.ln(3)
        pdf.cell(0, 9, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(*BLUE)
        pdf.set_line_width(0.6)
        pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
        pdf.ln(3)
        pdf.set_text_color(*BLACK)

    def h2(text):
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 11.5)
        pdf.set_text_color(*BLUE)
        pdf.ln(1.5)
        pdf.cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*BLACK)

    def paragraph(text, size=9.6):
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "", size)
        pdf.set_text_color(*BLACK)
        pdf.multi_cell(0, 4.9, text)
        pdf.set_x(pdf.l_margin)
        pdf.ln(0.5)

    def bullet(text, bold_lead, size=9.6):
        pdf.set_font("Helvetica", "", size)
        pdf.set_x(pdf.l_margin + 4)
        pdf.set_text_color(*BLACK)
        if bold_lead:
            pdf.set_font("Helvetica", "B", size)
            lead_w = pdf.get_string_width(bold_lead + " ")
            x0, y0 = pdf.get_x(), pdf.get_y()
            pdf.cell(lead_w, 4.9, bold_lead + " ")
            pdf.set_font("Helvetica", "", size)
            pdf.set_xy(x0 + lead_w, y0)
            pdf.multi_cell(pdf.w - pdf.r_margin - (x0 + lead_w), 4.9, text)
        else:
            pdf.multi_cell(text_w - 4, 4.9, f"-  {text}")
        pdf.ln(0.5)

    def callout(title, text, color):
        color = CALLOUT_COLORS[color]
        pdf.ln(1)
        x, y = pdf.l_margin, pdf.get_y()
        title_h, line_h = 6, 4.6
        avail_w = text_w - 8
        pdf.set_font("Helvetica", "", 9.3)
        lines = pdf.multi_cell(avail_w, line_h, text, dry_run=True, output="LINES")
        box_h = title_h + len(lines) * line_h + 6
        pdf.set_fill_color(*LGRAY)
        pdf.set_draw_color(*color)
        pdf.set_line_width(0.8)
        pdf.rect(x, y, text_w, box_h, style="DF")
        pdf.set_xy(x + 4, y + 3)
        pdf.set_font("Helvetica", "B", 9.8)
        pdf.set_text_color(*color)
        pdf.cell(0, title_h, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_xy(x + 4, pdf.get_y())
        pdf.set_font("Helvetica", "", 9.3)
        pdf.set_text_color(*BLACK)
        pdf.multi_cell(avail_w, line_h, text)
        pdf.set_xy(x, y + box_h + 4)

    def figure(filename, caption, width):
        pdf.set_x(pdf.l_margin)
        pdf.ln(2)
        if pdf.get_y() + width * 0.8 > pdf.h - 20:
            pdf.add_page()
        pdf.image(str(FIGURES_DIR / filename), x=(pdf.w - width) / 2, w=width)
        pdf.set_font("Helvetica", "I", 8.3)
        pdf.set_text_color(*GRAY)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 4.3, caption, align="C")
        pdf.set_text_color(*BLACK)
        pdf.set_x(pdf.l_margin)
        pdf.ln(2)

    def table(headers, rows):
        pdf.ln(1)
        # keep the whole table on one page (header + rows + notes)
        est_h = 7 + sum(6.5 + (4.2 if note else 0) for _, _, note in rows)
        if pdf.get_y() + est_h > pdf.h - pdf.b_margin:
            pdf.add_page()
        col_w = [100, 60]
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 9.5)
        pdf.set_fill_color(*NAVY)
        pdf.set_text_color(255, 255, 255)
        pdf.cell(col_w[0], 7, headers[0], fill=True)
        pdf.cell(col_w[1], 7, headers[1], fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*BLACK)
        pdf.set_fill_color(*ROW_FILL)
        for i, (name, val, note) in enumerate(rows):
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "B" if is_bold_row(name) else "", 9.3)
            fill = i % 2 == 0
            pdf.cell(col_w[0], 6.5, name, border="B", fill=fill)
            pdf.cell(col_w[1], 6.5, val, border="B", fill=fill, new_x="LMARGIN", new_y="NEXT")
            if note:
                pdf.set_font("Helvetica", "I", 8)
                pdf.set_text_color(*GRAY)
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 4.2, note)
                pdf.set_text_color(*BLACK)
        pdf.ln(2)

    # title page
    pdf.add_page()
    pdf.ln(50)
    pdf.set_font("Helvetica", "B", 23)
    pdf.set_text_color(*NAVY)
    pdf.cell(0, 11, TITLE["title"], align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 13)
    pdf.set_text_color(*GRAY)
    pdf.cell(0, 9, TITLE["subtitle"], align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(16)
    pdf.set_draw_color(*BLUE)
    pdf.set_line_width(0.8)
    pdf.line(pdf.w / 2 - 22, pdf.get_y(), pdf.w / 2 + 22, pdf.get_y())
    pdf.ln(16)
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(*BLACK)
    for line in TITLE["lines"]:
        pdf.cell(0, 7, line, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(30)
    pdf.set_font("Helvetica", "", 9.5)
    pdf.set_text_color(*GRAY)
    pdf.multi_cell(0, 5, TITLE["footer"], align="C")

    pdf.add_page()
    for kind, *args in blocks:
        if kind == "h1":
            h1(*args)
        elif kind == "h2":
            h2(*args)
        elif kind == "p":
            paragraph(*args)
        elif kind == "bullet":
            bullet(*args)
        elif kind == "callout":
            callout(*args)
        elif kind == "figure":
            figure(*args)
        elif kind == "table":
            table(*args)
        else:
            raise ValueError(f"unknown block type {kind!r}")

    pdf.output(str(path))
    print(f"wrote {path}")


# ---------------------------------------------------------------------------
# Word (python-docx)
# ---------------------------------------------------------------------------
def build_docx(blocks, logo, path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    def rgb(c):
        return RGBColor(*c)

    def hex_(c):
        return "{:02X}{:02X}{:02X}".format(*c)

    def set_border(p_or_cell_pr, side, color, size_eighths, tag="w:pBdr"):
        borders = p_or_cell_pr.find(qn(tag))
        if borders is None:
            borders = OxmlElement(tag)
            p_or_cell_pr.append(borders)
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size_eighths))
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), hex_(color))
        borders.append(el)

    def shade_cell(cell, color):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_(color))
        cell._tc.get_or_add_tcPr().append(shd)

    def add_page_field(paragraph):
        run = paragraph.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            run._r.append(el)
        return run

    doc = Document()

    # page setup: A4, same margins as the PDF, title page gets its own (logo-only) header
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(MARGIN_MM)
    section.top_margin, section.bottom_margin = Mm(22), Mm(18)
    section.header_distance = Mm(6)
    section.different_first_page_header_footer = True
    text_w = section.page_width - section.left_margin - section.right_margin

    # base styles
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (("Heading 1", 15, NAVY, 14, 6), ("Heading 2", 11.5, BLUE, 8, 3)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = rgb(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        # python-docx sets the latin font only; set the east-asian/theme font too so Word honors it
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), "Arial")
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)

    logo_bytes = None
    if logo is not None:
        buf = io.BytesIO()
        logo.save(buf, format="PNG")
        logo_bytes = buf.getvalue()

    # headers
    first = section.first_page_header.paragraphs[0]
    if logo_bytes:
        first.add_run().add_picture(io.BytesIO(logo_bytes), height=Mm(LOGO_H_MM))
    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(text_w, WD_TAB_ALIGNMENT.RIGHT)
    if logo_bytes:
        hp.add_run().add_picture(io.BytesIO(logo_bytes), height=Mm(LOGO_H_MM))
        hp.add_run("   ")
    r = hp.add_run(RUNNING_HEADER)
    r.font.size, r.font.italic, r.font.color.rgb = Pt(8), True, rgb(GRAY)
    hp.add_run("\t")
    pr = add_page_field(hp)
    pr.font.size, pr.font.color.rgb = Pt(8), rgb(GRAY)
    set_border(hp._p.get_or_add_pPr(), "bottom", LGRAY, 4)

    # title page
    def centered(text, size, color, bold=False, before=0, after=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        run = p.add_run(text)
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, rgb(color)
        return p

    centered(TITLE["title"], 23, NAVY, bold=True, before=150, after=4)
    centered(TITLE["subtitle"], 13, GRAY, after=30)
    rule = centered("", 4, BLUE, after=30)
    rule.paragraph_format.left_indent = rule.paragraph_format.right_indent = Mm(66)
    set_border(rule._p.get_or_add_pPr(), "bottom", BLUE, 12)
    for line in TITLE["lines"]:
        centered(line, 11, BLACK, after=2)
    last = centered(TITLE["footer"], 9.5, GRAY, before=80)
    last.add_run().add_break(WD_BREAK.PAGE)

    # body
    for kind, *args in blocks:
        if kind == "h1":
            p = doc.add_heading(args[0], level=1)
            set_border(p._p.get_or_add_pPr(), "bottom", BLUE, 12)
        elif kind == "h2":
            doc.add_heading(args[0], level=2)
        elif kind == "p":
            p = doc.add_paragraph(args[0])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "bullet":
            text, lead = args
            p = doc.add_paragraph(style="List Bullet")
            if lead:
                p.add_run(lead + " ").bold = True
            p.add_run(text)
        elif kind == "callout":
            title, text, color = args
            color = CALLOUT_COLORS[color]
            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = t.cell(0, 0)
            cell.width = text_w
            shade_cell(cell, LGRAY)
            tc_pr = cell._tc.get_or_add_tcPr()
            for side in ("top", "left", "bottom", "right"):
                set_border(tc_pr, side, color, 12, tag="w:tcBorders")
            # never split the box across pages
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            tp = cell.paragraphs[0]
            tp.paragraph_format.space_before = Pt(3)
            tp.paragraph_format.keep_with_next = True
            run = tp.add_run(title)
            run.bold, run.font.color.rgb = True, rgb(color)
            bp = cell.add_paragraph(text)
            bp.paragraph_format.space_after = Pt(4)
            bp.paragraph_format.keep_together = True
            for run in bp.runs:
                run.font.size = Pt(9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif kind == "figure":
            filename, caption, width = args
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(str(FIGURES_DIR / filename), width=Mm(width))
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            cp.paragraph_format.keep_together = True
            run = cp.add_run(caption)
            run.italic, run.font.size, run.font.color.rgb = True, Pt(8.5), rgb(GRAY)
        elif kind == "table":
            headers, rows = args
            col_w = [Mm(110), Mm(64)]
            t = doc.add_table(rows=1 + len(rows), cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            for j, h in enumerate(headers):
                cell = t.cell(0, j)
                cell.width = col_w[j]
                shade_cell(cell, NAVY)
                run = cell.paragraphs[0].add_run(h)
                run.bold, run.font.color.rgb = True, RGBColor(255, 255, 255)
            for i, (name, val, note) in enumerate(rows, start=1):
                bold = is_bold_row(name)
                for j, text in enumerate((name, val)):
                    cell = t.cell(i, j)
                    cell.width = col_w[j]
                    if i % 2 == 1:
                        shade_cell(cell, ROW_FILL)
                    set_border(cell._tc.get_or_add_tcPr(), "bottom", (200, 200, 200), 4, tag="w:tcBorders")
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(1)
                    cp.add_run(text).bold = bold
                    # keep rows together so the table doesn't split across pages
                    cp.paragraph_format.keep_with_next = i < len(rows)
                if note:
                    np_ = t.cell(i, 0).add_paragraph()
                    np_.paragraph_format.space_after = Pt(1)
                    np_.paragraph_format.keep_with_next = i < len(rows)
                    run = np_.add_run(note)
                    run.italic, run.font.size, run.font.color.rgb = True, Pt(8), rgb(GRAY)
            doc.add_paragraph()
        else:
            raise ValueError(f"unknown block type {kind!r}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    print(f"wrote {path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--pdf", action="store_true", help="build only the PDF")
    parser.add_argument("--docx", action="store_true", help="build only the Word document")
    args = parser.parse_args()
    both = not (args.pdf or args.docx)

    with open(METRICS_PATH) as f:
        blocks = build_blocks(json.load(f))
    logo = load_logo()

    if both or args.pdf:
        build_pdf(blocks, logo, PDF_PATH)
    if both or args.docx:
        build_docx(blocks, logo, DOCX_PATH)


if __name__ == "__main__":
    main()
